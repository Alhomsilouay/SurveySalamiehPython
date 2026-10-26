from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import os

src="Field surveying.xlsx"
out="Salamiyah_Urban_Information_Bank_Comprehensive.xlsx"
wb=load_workbook(src)

for s in ["README","01_Field_Survey","02_Raw_Data","03_Indicator_Dictionary","04_Lookups",
          "05_Calculations","06_Dashboard","07_GIS","08_Photos","09_Action_Plan"]:
    if s in wb.sheetnames: del wb[s]

blue=PatternFill("solid",fgColor="1F4E78")
light=PatternFill("solid",fgColor="D9EAF7")
white=Font(color="FFFFFF",bold=True)
bold=Font(bold=True)
thin=Side(style="thin",color="B7B7B7")
border=Border(left=thin,right=thin,top=thin,bottom=thin)
center=Alignment(horizontal="center",vertical="center",wrap_text=True)
wrap=Alignment(vertical="top",wrap_text=True)

# README
ws=wb.create_sheet("README",0)
for r in [
["بنك المعلومات الحضري لمدينة السلمية (SUIB)",""],
["الغرض","تحويل ملف المسح الميداني الحالي إلى قاعدة معلومات حضرية قابلة للتحديث والتحليل وGIS ولوحات المعلومات."],
["مصدر البناء","ملف المستخدم Field surveying.xlsx؛ تم الحفاظ على Info وData كمرجع أصلي."],
["المبدأ","الحفاظ على البيانات الخام وعدم اختزالها كلها إلى درجات. المؤشرات التقييمية لها مقياس 0–4، والقياسات الخام تبقى رقمية بوحداتها."],
["وحدة المسح","مقطع شارع/نقطة مسح عند اختلاف الظروف على امتداد الشارع."],
["المراحل","المسح الميداني → البيانات الخام → قاموس المؤشرات → الحسابات → Dashboard/GIS → خطة العمل."],
["الأوزان","لم تتم إضافة أوزان نهائية؛ تعتمد بعد تجربة ميدانية أولية واعتماد القاموس."]
]: ws.append(r)
for c in ws[1]: c.fill=blue;c.font=white;c.alignment=center
for row in ws.iter_rows():
    for c in row: c.border=border;c.alignment=wrap
ws.column_dimensions["A"].width=25;ws.column_dimensions["B"].width=110

# Lookups
lk=wb.create_sheet("04_Lookups")
groups={
"Street_Type":["رئيسي","ثانوي","تجاري","سكني","مختلط"],
"Traffic":["عالية","متوسطة","منخفضة"],
"Sensitivity":["عالية","متوسطة","منخفضة"],
"Sensitivity_Reason":["مدرسة","مشفى","سوق","مبنى حكومي","ساحة","موقع تراثي","جامع/موقع ديني","موقع سياحي","أخرى"],
"Score_0_4":["ممتاز","جيد","متوسط","سيئ","سيئ جداً"],
"Priority":["عالية","متوسطة","منخفضة"],
"Waste":["يوجد كثيرا","يوجد قليلا","لايوجد"]
}
for j,(k,vals) in enumerate(groups.items(),1):
    lk.cell(1,j,k);lk.cell(1,j).fill=blue;lk.cell(1,j).font=white;lk.cell(1,j).alignment=center
    for i,v in enumerate(vals,2): lk.cell(i,j,v)
    lk.column_dimensions[get_column_letter(j)].width=22

# Field survey
fs=wb.create_sheet("01_Field_Survey")
headers=[
"Record_ID","Street_Name","Sector","Neighborhood","Neighborhood_Tab","Street_No","Segment_ID","Street_Type",
"Traffic_Density","Sensitivity","Sensitivity_Reason","Survey_Date","GPS_Start","GPS_End","Photo_ID",
"Pavement_Condition","Damaged_Asphalt_m2","Sidewalk_Condition","Damaged_Sidewalk_m","Lighting_Quality",
"Facade_Condition","Advertisement_Condition","Murals_Public_Art","Aesthetic_Lighting","Street_Cleanliness",
"Bins_240L_Required","Large_Bins_Required","Air_Odor_Quality","Construction_Waste","Pedestrian_Safety",
"Accessibility","Noise_Level","Night_Safety","Bike_Lanes","Tree_Density","Landscaping_Quality","Irrigation",
"Shade","Parks","Squares_Roundabouts","Field_Priority","Notes"]
ar=[
"المتسلسل","اسم الشارع","القطاع","الحي","تبويب الحي","رقم الشارع","معرف المقطع","نوع الشارع","كثافة الحركة",
"حساسية الموقع","سبب الحساسية","تاريخ المسح","GPS البداية","GPS النهاية","معرف الصورة","حالة الرصف",
"مساحة الإسفلت المتضررة (م²)","حالة الأرصفة","طول الأرصفة المتضررة (م)","جودة الإنارة","حالة الواجهات",
"حالة الإعلانات","الجداريات والفن العام","الإضاءة الجمالية","نظافة الشارع","حاويات 240 لتر المطلوبة",
"الحاويات الكبيرة المطلوبة","جودة الهواء والروائح","مخلفات بناء ونفايات","سلامة المشاة","إمكانية الوصول",
"الهدوء وقلة الضوضاء","الأمان الليلي","مسارات الدراجات","كثافة الأشجار","جودة التشجير","الري","الظلال",
"الحدائق","الساحات والدورات","أولوية التدخل الميدانية","ملاحظات"]
fs.append(["استمارة المسح الميداني – بنك المعلومات الحضري لمدينة السلمية"]+[""]*(len(headers)-1))
fs.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
fs["A1"].fill=blue;fs["A1"].font=Font(color="FFFFFF",bold=True,size=14);fs["A1"].alignment=center
fs.append(headers);fs.append(ar)
for rr in [2,3]:
    for c in fs[rr]:
        c.fill=blue if rr==2 else light;c.font=white if rr==2 else bold;c.alignment=center;c.border=border

# Transfer user's current Data sheet
srcws=wb["Data"]
for r in range(4,srcws.max_row+1):
    vals=[srcws.cell(r,c).value for c in range(1,srcws.max_column+1)]
    if not any(v not in (None,"") for v in vals): continue
    outrow=[None]*len(headers)
    mapping={
      1:1,2:2,3:3,4:4,5:5,6:6,8:7,7:8,12:9,11:10,10:15,9:12,
      13:16,14:17,15:18,16:19,17:20,19:21,20:22,21:23,22:24,
      23:25,24:26,25:27,26:28,27:29,28:30,29:31,30:32,31:33,32:34,
      33:35,34:36,35:37,36:38,37:39,38:40,39:41,40:42}
    for target,source in mapping.items():
        if source<=len(vals): outrow[target-1]=vals[source-1]
    outrow[6]=f"S{r-3:04d}"
    fs.append(outrow)

fs.freeze_panes="A4"
for i in range(1,len(headers)+1): fs.column_dimensions[get_column_letter(i)].width=17
for i in [2,4,5,11,42]: fs.column_dimensions[get_column_letter(i)].width=23
for row in fs.iter_rows(min_row=4):
    for c in row: c.border=border;c.alignment=wrap

def add_dv(col, lookup):
    dv=DataValidation(type="list",formula1=f"'04_Lookups'!${lookup}$2:${lookup}$20",allow_blank=True)
    fs.add_data_validation(dv);dv.add(f"{get_column_letter(col)}4:{get_column_letter(col)}1000")
add_dv(8,"A");add_dv(9,"B");add_dv(10,"C");add_dv(11,"D")
for c in [16,18,20,21,22,23,24,25,28,30,31,32,33,34,35,36,37,38,39,40]: add_dv(c,"E")
add_dv(29,"G");add_dv(41,"F")

# Raw data
raw=wb.create_sheet("02_Raw_Data")
raw_headers=["Record_ID","Street_Name","Segment_ID","Survey_Date","GPS_Start","GPS_End","Damaged_Asphalt_m2",
"Damaged_Sidewalk_m","Bins_240L_Required","Large_Bins_Required","Tree_Count","Healthy_Tree_Count","Dry_Tree_Count",
"Green_Area_m2","Planted_Length_m","Park_Count","Bench_Count","Irrigation_Point_Count","Lighting_Units",
"Working_Lighting_Units","Notes"]
raw.append(raw_headers)
raw.append(["معرف السجل","اسم الشارع","معرف المقطع","تاريخ المسح","GPS البداية","GPS النهاية","مساحة الإسفلت المتضررة (م²)",
"طول الأرصفة المتضررة (م)","حاويات 240 لتر المطلوبة","الحاويات الكبيرة المطلوبة","عدد الأشجار","عدد الأشجار السليمة",
"عدد الأشجار الجافة","المساحة الخضراء (م²)","طول التشجير (م)","عدد الحدائق","عدد المقاعد","عدد نقاط الري",
"عدد وحدات الإنارة","عدد وحدات الإنارة العاملة","ملاحظات"])
for rr in [1,2]:
    for c in raw[rr]: c.fill=blue if rr==1 else light;c.font=white if rr==1 else bold;c.alignment=center;c.border=border
for r in range(3,1003):
    pairs={1:1,2:2,3:7,4:12,5:13,6:14,7:17,8:19,9:26,10:27}
    for dest,src_col in pairs.items():
        raw.cell(r,dest,f'=IF(\'01_Field_Survey\'!{get_column_letter(src_col)}{r}="","",\'01_Field_Survey\'!{get_column_letter(src_col)}{r})')
for i in range(1,len(raw_headers)+1): raw.column_dimensions[get_column_letter(i)].width=20

# Dictionary
dic=wb.create_sheet("03_Indicator_Dictionary")
dh=["ID","المحور","المؤشر","نوع البيانات","التعريف","طريقة القياس","الوحدة/المقياس","قاعدة الدرجة","بيانات خام مطلوبة","صورة","التكرار","ملاحظات"]
dic.append(dh)
for c in dic[1]: c.fill=blue;c.font=white;c.alignment=center;c.border=border
items=[
("GEO-01","GIS","نوع الشارع","تصنيف","التصنيف الوظيفي للشارع","قائمة","رئيسي/ثانوي/تجاري/سكني/مختلط","لا يوجد","لا","اختياري","عند المسح",""),
("GEO-02","GIS","كثافة الحركة","تصنيف","مستوى حركة المركبات والمشاة","تقييم ميداني","عالية/متوسطة/منخفضة","لا يوجد","يفضل عدّ لاحقاً","اختياري","عند المسح",""),
("GEO-03","GIS","حساسية الموقع","تصنيف","أهمية الموقع أو حساسيته","تصنيف","عالية/متوسطة/منخفضة","لا يوجد","سبب الحساسية","نعم","عند المسح","يفصل عن سبب الحساسية"),
("INF-01","البنية التحتية","حالة الرصف","تقييمي","الحالة العامة لسطح الطريق","فحص بصري","0–4","4 ممتاز، 3 جيد، 2 متوسط، 1 سيئ، 0 سيئ جداً","مساحة الضرر","نعم","سنوي/عند الحاجة",""),
("INF-02","البنية التحتية","مساحة الإسفلت المتضررة","خام","مساحة الجزء المتضرر","قياس/تقدير","م²","لا يوجد","نعم","نعم","عند المسح",""),
("INF-03","البنية التحتية","حالة الأرصفة","تقييمي","الحالة العامة للأرصفة","فحص بصري","0–4","4 ممتاز، 3 جيد، 2 متوسط، 1 سيئ، 0 معدوم","طول الضرر","نعم","سنوي/عند الحاجة",""),
("INF-04","البنية التحتية","طول الأرصفة المتضررة","خام","طول الجزء المتضرر","قياس","م","لا يوجد","نعم","نعم","عند المسح",""),
("INF-05","البنية التحتية","جودة الإنارة","تقييمي","كفاءة الإنارة الأساسية","فحص بصري/ليلي","0–4","4 ممتاز ... 0 سيئ جداً","عدد وحدات الإنارة","نعم","سنوي","تفصل عن الإضاءة الجمالية"),
("VIS-01","الهوية البصرية","حالة الواجهات","تقييمي","الحالة البصرية والمادية للواجهات","فحص بصري","0–4","4 ممتاز ... 0 سيئ جداً","لا","نعم","سنوي",""),
("VIS-02","الهوية البصرية","حالة الإعلانات","تقييمي","تنظيم وحالة الإعلانات","فحص بصري","0–4","4 ممتاز ... 0 سيئ جداً","لا","نعم","سنوي",""),
("VIS-03","الهوية البصرية","الجداريات والفن العام","تقييمي","وجود وحالة الأعمال الفنية العامة","فحص بصري","0–4","4 ممتاز ... 0 سيئ جداً","لا","نعم","سنوي",""),
("VIS-04","الهوية البصرية","الإضاءة الجمالية","تقييمي","الإضاءة المخصصة للتجميل","فحص ليلي","0–4","4 ممتاز ... 0 سيئ جداً","لا","نعم","موسمي/سنوي",""),
("CLN-01","النظافة","نظافة الشارع","تقييمي","مستوى النظافة العامة","فحص بصري","0–4","4 ممتاز ... 0 سيئ جداً","لا","نعم","دوري",""),
("CLN-02","النظافة","حاويات 240 لتر المطلوبة","خام","عدد الحاويات المطلوبة","عدّ/تقدير","عدد","لا يوجد","نعم","نعم","دوري",""),
("CLN-03","النظافة","الحاويات الكبيرة المطلوبة","خام","عدد الحاويات الكبيرة المطلوبة","عدّ/تقدير","عدد","لا يوجد","نعم","نعم","دوري",""),
("CLN-04","النظافة","جودة الهواء والروائح","تقييمي","مستوى الروائح/جودة الهواء ميدانياً","ملاحظة/قياس","0–4","4 ممتاز ... 0 سيئ جداً","اختياري","نعم","دوري",""),
("CLN-05","النظافة","مخلفات بناء ونفايات","تقييمي","وجود وكثافة المخلفات","ملاحظة ميدانية","0–4","4 لا توجد، 3 قليلة، 2 متوسطة، 1 كثيرة، 0 كثيرة جداً","لا","نعم","دوري",""),
("URB-01","الحضر","سلامة المشاة","تقييمي","سلامة الحركة للمشاة","فحص ميداني","0–4","4 ممتاز ... 0 سيئ جداً","لا","نعم","سنوي",""),
("URB-02","الحضر","إمكانية الوصول","تقييمي","ملاءمة المكان للوصول الشامل","فحص الأرصفة والمعابر","0–4","4 ممتاز ... 0 سيئ جداً","لا","نعم","سنوي",""),
("URB-03","الحضر","الهدوء وقلة الضوضاء","تقييمي","مستوى الهدوء النسبي","ملاحظة/قياس dB","0–4","4 هادئ جداً ... 0 شديد الضوضاء","يفضل dB خام","نعم","دوري",""),
("URB-04","الحضر","الأمان الليلي","تقييمي","مؤشرات الأمان الميدانية ليلاً","فحص ليلي","0–4","4 ممتاز ... 0 سيئ جداً","لا","نعم","سنوي",""),
("URB-05","الحضر","مسارات الدراجات","تقييمي","توفر وجودة البنية للدراجات","فحص ميداني","0–4","4 ممتاز ... 0 معدوم","طول المسار لاحقاً","نعم","سنوي",""),
("GRN-01","المساحات الخضراء","كثافة الأشجار","تقييمي","كثافة الأشجار في المقطع","تقييم/عدّ","0–4","4 مرتفعة جداً ... 0 معدومة","عدد الأشجار","نعم","سنوي",""),
("GRN-02","المساحات الخضراء","جودة التشجير","تقييمي","حالة وجودة التشجير","فحص بصري","0–4","4 ممتاز ... 0 سيئ جداً","عدد السليم/الجاف","نعم","موسمي",""),
("GRN-03","المساحات الخضراء","الري","تقييمي","كفاءة وتوفر الري","فحص ميداني","0–4","4 ممتاز ... 0 معدوم","نقاط الري","نعم","موسمي",""),
("GRN-04","المساحات الخضراء","الظلال","تقييمي","توفر الظل","ملاحظة ميدانية","0–4","4 مرتفع جداً ... 0 معدوم","اختياري","نعم","موسمي",""),
("GRN-05","المساحات الخضراء","الحدائق","تقييمي","وجود وحالة الحدائق","فحص ميداني","0–4","4 ممتاز ... 0 معدوم","عدد الحدائق","نعم","سنوي",""),
("GRN-06","المساحات الخضراء","الساحات والدورات","تقييمي","وجود وحالة الساحات والدورات","فحص ميداني","0–4","4 ممتاز ... 0 معدوم","المساحة لاحقاً","نعم","سنوي",""),
("ACT-01","الإدارة","أولوية التدخل الميدانية","تصنيف","أولوية أولية يحددها الفريق","تقييم الفريق","عالية/متوسطة/منخفضة","لا يوجد","بيانات المؤشرات","اختياري","عند المسح","تضاف لاحقاً أولوية محسوبة منفصلة")
]
for x in items: dic.append(x)
for row in dic.iter_rows(min_row=2):
    for c in row:c.border=border;c.alignment=wrap
for i,w in enumerate([14,22,28,16,42,35,24,45,24,12,18,40],1): dic.column_dimensions[get_column_letter(i)].width=w

# Calculations
calc=wb.create_sheet("05_Calculations")
for r in [
["عنصر الحساب","المنطق","الهدف"],
["عدد سجلات المسح","=COUNTA('01_Field_Survey'!A4:A1000)","حجم قاعدة البيانات"],
["إجمالي مساحة الإسفلت المتضررة","=SUM('01_Field_Survey'!Q4:Q1000)","بيان خام"],
["إجمالي طول الأرصفة المتضررة","=SUM('01_Field_Survey'!S4:S1000)","بيان خام"],
["إجمالي حاويات 240 لتر المطلوبة","=SUM('01_Field_Survey'!Z4:Z1000)","بيان خام"],
["إجمالي الحاويات الكبيرة المطلوبة","=SUM('01_Field_Survey'!AA4:AA1000)","بيان خام"],
["الأولوية المحسوبة","تُبنى لاحقاً من الحالة + الحساسية + كثافة الحركة + حجم الضرر","تقليل الذاتية"],
["مبدأ","لا تستبدل البيانات الخام بالدرجات","إبقاء القاعدة قابلة لإعادة التحليل"]
]: calc.append(r)
for c in calc[1]: c.fill=blue;c.font=white;c.alignment=center;c.border=border
for row in calc.iter_rows():
    for c in row:c.border=border;c.alignment=wrap
calc.column_dimensions["A"].width=34;calc.column_dimensions["B"].width=75;calc.column_dimensions["C"].width=45

# Dashboard
dash=wb.create_sheet("06_Dashboard")
for r in [
["لوحة المعلومات الحضرية – نسخة أولية",""],
["المؤشر","القيمة"],
["عدد سجلات المسح","=COUNTA('01_Field_Survey'!A4:A1000)"],
["إجمالي مساحة الإسفلت المتضررة (م²)","=SUM('01_Field_Survey'!Q4:Q1000)"],
["إجمالي طول الأرصفة المتضررة (م)","=SUM('01_Field_Survey'!S4:S1000)"],
["إجمالي حاويات 240 لتر المطلوبة","=SUM('01_Field_Survey'!Z4:Z1000)"],
["إجمالي الحاويات الكبيرة المطلوبة","=SUM('01_Field_Survey'!AA4:AA1000)"],
["ملاحظة","لوحة بداية؛ بعد اعتماد القاموس والأوزان يمكن إضافة Pivot Tables وخرائط حرارية ومؤشرات حسب الحي والقطاع."]
]: dash.append(r)
for c in dash[1]: c.fill=blue;c.font=white;c.alignment=center;c.border=border
for row in dash.iter_rows():
    for c in row:c.border=border;c.alignment=wrap
dash.column_dimensions["A"].width=42;dash.column_dimensions["B"].width=80

# GIS
gis=wb.create_sheet("07_GIS")
for r in [
["حقل","المحتوى","الغرض"],["Record_ID","معرف السجل","مفتاح الربط"],["Street_Name","اسم الشارع",""],
["Segment_ID","معرف المقطع","يجب أن يكون فريداً"],["GPS_Start","إحداثيات البداية","GIS"],["GPS_End","إحداثيات النهاية","GIS"],
["Sensitivity","حساسية الموقع","طبقة تحليلية"],["Calculated_Priority","الأولوية المحسوبة","بعد اعتماد المعادلة"],["Photo_ID","معرف الصورة","ربط الصور"]
]: gis.append(r)
for c in gis[1]: c.fill=blue;c.font=white;c.alignment=center;c.border=border
for row in gis.iter_rows():
    for c in row:c.border=border;c.alignment=wrap
gis.column_dimensions["A"].width=25;gis.column_dimensions["B"].width=50;gis.column_dimensions["C"].width=55

# Photos
ph=wb.create_sheet("08_Photos")
ph.append(["Photo_ID","Record_ID","Street_Name","Segment_ID","Photo_Type","File_Name_or_Link","Date","Description"])
ph.append(["معرف الصورة","معرف السجل","اسم الشارع","معرف المقطع","نوع الصورة","اسم الملف/الرابط","التاريخ","وصف مختصر"])
for rr in [1,2]:
    for c in ph[rr]: c.fill=blue if rr==1 else light;c.font=white if rr==1 else bold;c.alignment=center;c.border=border
for i,w in enumerate([18,18,25,20,22,45,16,50],1): ph.column_dimensions[get_column_letter(i)].width=w

# Action plan
act=wb.create_sheet("09_Action_Plan")
act.append(["Action_ID","Record_ID","Street_Name","Issue","Evidence","Proposed_Action","Responsible_Entity","Estimated_Cost","Urgency","Status","Target_Date","Completion_Date","Notes"])
act.append(["معرف الإجراء","معرف السجل","اسم الشارع","المشكلة","الدليل/الصورة","الإجراء المقترح","الجهة المسؤولة","التكلفة التقديرية","الاستعجال","الحالة","التاريخ المستهدف","تاريخ الإنجاز","ملاحظات"])
for rr in [1,2]:
    for c in act[rr]: c.fill=blue if rr==1 else light;c.font=white if rr==1 else bold;c.alignment=center;c.border=border
for i,w in enumerate([16,18,25,30,30,45,28,20,18,18,18,18,45],1): act.column_dimensions[get_column_letter(i)].width=w

for sh in wb.worksheets:
    sh.sheet_view.showGridLines=False
    sh.page_setup.fitToWidth=1
    sh.page_setup.fitToHeight=0
    sh.sheet_properties.pageSetUpPr.fitToPage=True

wb.save(out)
print(out)
print(os.path.getsize(out))
